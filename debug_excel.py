import openpyxl
from openpyxl.cell import MergedCell

# テンプレートファイルを読み込み
wb = openpyxl.load_workbook('degital_value_list.xlsx')
ws = wb.active

print(f"ワークシート名: {ws.title}")
print(f"最大行: {ws.max_row}")
print(f"最大列: {ws.max_column}")

# マージされたセルの範囲を確認
print("\nマージされたセル範囲:")
for merged_range in ws.merged_cells.ranges:
    print(f"  {merged_range}")

# 各セルの状態を確認（最初の10行×14列）
print("\nセルの状態確認:")
for row in range(1, min(11, ws.max_row + 1)):
    for col in range(1, 15):
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, MergedCell):
            print(f"  行{row}列{col}: MergedCell")
        else:
            print(f"  行{row}列{col}: 通常セル (値: {cell.value})")