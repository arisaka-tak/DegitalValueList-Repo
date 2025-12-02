# Generated manually

from django.db import migrations, models
import django.db.models.deletion
from openpyxl import load_workbook
import os

def create_manufacturer_master(apps, schema_editor):
    """Excelファイルからメーカーマスタを作成"""
    Manufacturer = apps.get_model('products_master', 'Manufacturer')
    Product = apps.get_model('products_master', 'Product')
    
    # Excelファイルからメーカーリストを読み込み
    excel_path = r'C:\Project\ZCS_DegitalValueList\maker_list.xlsx'
    if os.path.exists(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
        
        manufacturers = set()
        for row in range(2, ws.max_row + 1):  # ヘッダーをスキップ
            cell_value = ws.cell(row=row, column=1).value
            if cell_value and cell_value.strip():
                manufacturers.add(cell_value.strip())
        
        # メーカーマスタに登録
        for manufacturer_name in sorted(manufacturers):
            Manufacturer.objects.get_or_create(name=manufacturer_name)
    
    # 既存商品のメーカー名からもマスタを作成
    existing_manufacturers = Product.objects.exclude(
        manufacturer__isnull=True
    ).exclude(
        manufacturer__exact=''
    ).values_list('manufacturer', flat=True).distinct()
    
    for manufacturer_name in existing_manufacturers:
        if manufacturer_name and manufacturer_name.strip():
            Manufacturer.objects.get_or_create(name=manufacturer_name.strip())

def convert_manufacturer_to_fk(apps, schema_editor):
    """既存商品のメーカー文字列を外部キーに変換"""
    Product = apps.get_model('products_master', 'Product')
    Manufacturer = apps.get_model('products_master', 'Manufacturer')
    
    # 既存データを変換
    for product in Product.objects.all():
        if hasattr(product, 'manufacturer') and product.manufacturer:
            try:
                manufacturer_obj = Manufacturer.objects.get(name=product.manufacturer)
                # 一時的にmanufacturer_fk_idフィールドに値を設定
                from django.db import connection
                with connection.cursor() as cursor:
                    cursor.execute(
                        "UPDATE products_master_product SET manufacturer_fk_id = %s WHERE id = %s",
                        [manufacturer_obj.id, product.id]
                    )
            except Manufacturer.DoesNotExist:
                pass

def reverse_manufacturer_master(apps, schema_editor):
    """逆マイグレーション用"""
    pass

class Migration(migrations.Migration):

    dependencies = [
        ('products_master', '0025_add_master_tables'),
    ]

    operations = [
        migrations.CreateModel(
            name='Manufacturer',
            fields=[
                ('id', models.AutoField(auto_created=True, primary_key=True, serialize=False, verbose_name='ID')),
                ('name', models.CharField(max_length=200, unique=True, verbose_name='メーカー名')),
                ('is_active', models.BooleanField(default=True, verbose_name='有効')),
                ('created_at', models.DateTimeField(auto_now_add=True, verbose_name='作成日時')),
                ('updated_at', models.DateTimeField(auto_now=True, verbose_name='更新日時')),
            ],
            options={
                'verbose_name': 'メーカーマスタ',
                'verbose_name_plural': 'メーカーマスタ',
                'ordering': ['name'],
            },
        ),
        migrations.RunPython(create_manufacturer_master, reverse_manufacturer_master),
        migrations.AddField(
            model_name='product',
            name='manufacturer_fk',
            field=models.ForeignKey(blank=True, null=True, on_delete=django.db.models.deletion.SET_NULL, to='products_master.manufacturer', verbose_name='メーカーFK'),
        ),
        migrations.RunPython(convert_manufacturer_to_fk, reverse_manufacturer_master),
        migrations.RemoveField(
            model_name='product',
            name='manufacturer',
        ),
        migrations.RenameField(
            model_name='product',
            old_name='manufacturer_fk',
            new_name='manufacturer',
        ),
    ]